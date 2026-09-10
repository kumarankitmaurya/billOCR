"""Turns one supplier's ingested bills into the per-company book-ledger workbook.

Layout (matches the shop's existing book, per output-format.md):
  workbook = supplier
  sheet    = company
  each bill is a stacked, dated block inside its company's sheet:

      DJ-STI-36600   2025-08-17
      JAI MATA DI SRT   pc    price   tax    margin   SP  <- label row, once per sheet
      SAGAR             4     595     0      119      714
      SWIGGY            4     595
      (blank separator)
      DJ-STI-36742   2025-10-02                        <- next bill appends below
      DIWALI SPL        6     720

Columns D (tax) and E (margin) are *amounts*, not the raw percentages the
review screen captured (tax_pct, margin_pct) — computed here, independently,
for the book's benefit:
  D = rate * tax_pct / 100
  E = (rate + D) * margin_pct / 100   (margin on the tax-inclusive price,
                                        matching the review screen's chain)
Column F (SP / final price) is `final_price` exactly as it came from the
review screen — never recomputed from D/E here, since the frontend already
lets it be hand-overridden independently of tax_pct/margin_pct (see
app/models.py). Any of the three is left blank if its underlying field
wasn't set — these are shop-internal pricing decisions never read off the
bill itself.
"""

import io
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app import db
from app.config import settings

_UNSAFE_FILENAME_CHARS = re.compile(r'[\\/:*?"<>|]')


def _autofit(sheet: Worksheet) -> None:
    """Widen each column to roughly fit its longest cell value."""
    for col_cells in sheet.columns:
        length = max((len(str(cell.value)) for cell in col_cells if cell.value is not None), default=0)
        sheet.column_dimensions[get_column_letter(col_cells[0].column)].width = max(10, length + 2)


def _safe_sheet_title(name: str, used: set[str]) -> str:
    """Excel sheet titles must be <=31 chars and unique within the workbook."""
    title = name[:31]
    original = title
    counter = 2
    while title in used:
        title = f"{original[:28]}-{counter}"
        counter += 1
    used.add(title)
    return title


def _write_company_sheet(
    sheet: Worksheet,
    company_name: str,
    bills: list[tuple[str, str, list[tuple[str, int, float, float | None, float | None, float | None]]]],
) -> None:
    label_written = False
    for bill_no, bill_date, lines in bills:
        sheet.append([bill_no, bill_date])
        sheet.cell(row=sheet.max_row, column=1).font = Font(bold=True)

        if not label_written:
            sheet.append([company_name, "pc", "price", "tax", "margin", "SP"])
            sheet.cell(row=sheet.max_row, column=1).font = Font(italic=True)
            label_written = True

        for product, pcs, rate, final_price, margin_pct, tax_pct in lines:
            tax_amount = round(rate * tax_pct / 100, 2) if tax_pct is not None else None
            margin_amount = round((rate + (tax_amount or 0)) * margin_pct / 100, 2) if margin_pct is not None else None
            sheet.append([product, pcs, rate, tax_amount, margin_amount, final_price])

        sheet.append([])  # blank separator between bills

    _autofit(sheet)


def _output_path(supplier_name: str) -> Path:
    """Where this supplier's workbook is saved on disk (see settings.output_dir)."""
    safe_name = _UNSAFE_FILENAME_CHARS.sub("_", supplier_name).strip() or "supplier"
    return Path(settings.output_dir) / f"{safe_name}.xlsx"


def build_supplier_workbook(supplier_name: str) -> io.BytesIO:
    """Build the supplier's full book from the DB, save it under
    settings.output_dir, and return it as an in-memory buffer.

    Raises KeyError if the supplier has no ingested bills.
    """
    data = db.get_workbook_data(supplier_name)

    workbook = Workbook()
    workbook.remove(workbook.active)

    used_titles: set[str] = set()
    for company_name, bills in data.items():
        sheet = workbook.create_sheet(_safe_sheet_title(company_name, used_titles))
        _write_company_sheet(sheet, company_name, bills)

    workbook.save(_output_path(supplier_name))

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer
