"""The staff path end to end: a scan held as a draft, reviewed without the
rate, saved, and the doubtful lines left for the owner.

The scan is put straight into the draft table rather than through /preview,
which would spend a real OCR call.
"""

import io

from openpyxl import load_workbook

from app import db
from app.models import Article, BillExtraction, ExtractionResult

from .conftest import make_result

COST_KEYS = {"rate", "amount", "tax_pct", "margin_pct"}


def stored_scan(articles=None) -> str:
    articles = articles or [
        Article(company="JAI MATA DI SRT", product="GREEN TEA", pcs=4, rate=567.0,
                amount=2268.0, tax_pct=5.0, margin_pct=15.0, final_price=690.0),
        Article(company="JAI MATA DI SRT", product="SAGAR", pcs=4, rate=595.0,
                amount=9999.0, tax_pct=5.0, margin_pct=15.0, final_price=720.0),
    ]
    bill = BillExtraction(supplier="Dindayal Jalan", bill_no="DJ-1", bill_date="2025-08-17",
                          articles=articles)
    return db.save_draft(ExtractionResult(source_filename="bill.jpg", engine="groq", bill=bill))


def review(draft_id, rows, **bill):
    return {
        "supplier": "Dindayal Jalan",
        "drafts": [{
            "draft_id": draft_id,
            "bill_no": bill.get("bill_no", "DJ-1"),
            "bill_date": bill.get("bill_date", "2025-08-17"),
            "articles": rows,
        }],
    }


def row(ref, product, final_price, pcs=4, **extra):
    return {"ref": ref, "company": "JAI MATA DI SRT", "product": product, "pcs": pcs,
            "final_price": final_price, **extra}


def test_staff_save_keeps_the_stored_rate_and_marks_doubtful_lines(client, owner, supplier):
    draft_id = stored_scan()
    response = client.post("/api/bills/ingest", json=review(draft_id, [
        row(0, "GREEN TEA", 700.0, rate=1.0),   # rate from staff is ignored
        row(1, "SAGAR", 720.0),                 # amount mismatch -> owner
        row(None, "NEW ROW", 500.0, pcs=2),     # hand-added -> owner
    ]))
    assert response.status_code == 200, response.text
    assert response.json()["needs_check"] == 2

    checks = owner.get("/api/bills/checks").json()
    assert [c["product"] for c in checks] == ["SAGAR", "NEW ROW"]
    assert checks[0]["rate"] == 595.0
    assert checks[1]["rate"] is None

    found = owner.get("/api/bills/search", params={"name": "GREEN"}).json()
    assert found[0]["rate"] == 567.0
    assert found[0]["final_price"] == 700.0


def test_staff_cannot_ingest_whole_results_or_see_checks(client, supplier):
    assert client.post("/api/bills/ingest", json={"results": [make_result()]}).status_code == 403
    assert client.get("/api/bills/checks").status_code == 401
    assert client.patch("/api/bills/lines/1", json={"rate": 1.0}).status_code == 401


def test_an_unknown_draft_is_410(client, supplier):
    for draft_id in ("00000000-0000-0000-0000-000000000000", "not-a-uuid"):
        response = client.post("/api/bills/ingest", json=review(draft_id, []))
        assert response.status_code == 410


def test_resolving_a_check_clears_it_only_once_there_is_a_rate(client, owner, supplier):
    client.post("/api/bills/ingest", json=review(stored_scan(), [
        row(0, "GREEN TEA", 700.0), row(None, "NEW ROW", 500.0, pcs=2),
    ]))
    [line] = owner.get("/api/bills/checks").json()

    still = owner.patch(f"/api/bills/lines/{line['id']}", json={"final_price": 520.0})
    assert still.json() == {"id": line["id"], "needs_check": True}

    done = owner.patch(f"/api/bills/lines/{line['id']}", json={"rate": 400.0})
    assert done.json() == {"id": line["id"], "needs_check": False}
    assert owner.get("/api/bills/checks").json() == []

    assert owner.patch("/api/bills/lines/999999", json={"rate": 1.0}).status_code == 404


def test_staff_workbook_and_pricing_carry_no_cost(client, owner, supplier):
    client.post("/api/bills/ingest", json=review(stored_scan(), [row(0, "GREEN TEA", 700.0)]))

    staff_rows = list(load_workbook(io.BytesIO(
        client.get("/api/bills/workbook", params={"supplier": "Dindayal Jalan"}).content
    ))["JAI MATA DI SRT"].iter_rows(values_only=True))
    assert ("GREEN TEA", 4, 700) in staff_rows
    assert all(567 not in r for r in staff_rows)

    owner_rows = list(load_workbook(io.BytesIO(
        owner.get("/api/bills/workbook", params={"supplier": "Dindayal Jalan"}).content
    ))["JAI MATA DI SRT"].iter_rows(values_only=True))
    assert any(r[:3] == ("GREEN TEA", 4, 567) for r in owner_rows)

    assert client.get("/api/bills/pricing").json() == {"price_step": 5}
    assert "margin_pct_below" in owner.get("/api/bills/pricing").json()


def test_margin_rules_round_trip(client):
    db.set_margin_rule("Rohan  fab srt", "milk cake", 22.0)
    db.set_margin_rule("ROHAN FAB SRT", None, 19.0)
    assert db.margin_rules() == {("ROHAN FAB SRT", "MILK CAKE"): 22.0, ("ROHAN FAB SRT", ""): 19.0}
    assert db.remove_margin_rule("rohan fab srt", "MILK CAKE") is True
    assert db.remove_margin_rule("rohan fab srt", "MILK CAKE") is False


def test_preview_returns_no_cost_to_staff_and_a_draft_to_save(client, owner, supplier, monkeypatch):
    """Through the real endpoint, with only the OCR call stubbed out."""
    from app.routers import bills
    from app.services import ocr_strategy

    def fake_extract(**kwargs):
        bill = BillExtraction(supplier=kwargs["supplier"], bill_no="DJ-7", bill_date="2025-09-01",
                              articles=[Article(company="JAI MATA DI SRT", product="GREEN TEA",
                                                pcs=4, rate=567.0, amount=2268.0)])
        ocr_strategy._apply_pricing_defaults(bill, kwargs["margin_rules"])
        return ExtractionResult(source_filename="bill.jpg", engine="groq", bill=bill)

    monkeypatch.setattr(bills, "extract", fake_extract)
    db.set_margin_rule("JAI MATA DI SRT", "GREEN TEA", 30.0)

    jpeg = b"\xff\xd8\xff\xe0" + b"\x00" * 64
    upload = {"files": ("bill.jpg", jpeg, "image/jpeg")}
    form = {"supplier": "Dindayal Jalan"}

    staff = client.post("/api/bills/preview", files=upload, data=form)
    assert staff.status_code == 200, staff.text
    [view] = staff.json()

    def keys_in(value):
        if isinstance(value, dict):
            return set(value) | {k for v in value.values() for k in keys_in(v)}
        if isinstance(value, list):
            return {k for v in value for k in keys_in(v)}
        return set()

    assert not keys_in(view) & COST_KEYS
    # 567 x 1.05 x 1.30 = 773.96 -> rounded up to 775, from the product's rule
    assert view["bill"]["articles"][0]["final_price"] == 775.0

    [full] = owner.post("/api/bills/preview", files=upload, data=form).json()
    assert full["bill"]["articles"][0]["rate"] == 567.0
    assert full["bill"]["articles"][0]["margin_pct"] == 30.0

    saved = client.post("/api/bills/ingest", json=review(
        view["draft_id"], [row(0, "GREEN TEA", 775.0)], bill_no="DJ-7", bill_date="2025-09-01",
    ))
    assert saved.json()["needs_check"] == 0
