"""The book of record, over HTTP, against a real Postgres.

Covers what the SQLite-to-Postgres port could plausibly break: ingest and
its idempotency, search (including the case-insensitivity that SQLite gave
for free and Postgres doesn't), the admin gate on base price, and workbook
generation, plus the app-access gate and search pagination. /preview and
/extract aren't here — they call out to an LLM.
"""

import io

import pytest
from openpyxl import load_workbook

from .conftest import ADMIN_PASSWORD, make_result

ADMIN = {"X-Admin-Password": ADMIN_PASSWORD}


def test_ingest_persists_and_is_idempotent(client):
    payload = {"results": [make_result()], "supplier": "Dindayal Jalan"}

    first = client.post("/api/bills/ingest", json=payload)
    assert first.status_code == 200
    assert first.json() == {"ingested": 1, "suppliers": ["Dindayal Jalan"]}

    # Re-ingesting the same (supplier, bill_no) replaces its lines. If the
    # ON CONFLICT upsert regressed, this would either duplicate or 500.
    again = client.post("/api/bills/ingest", json=payload)
    assert again.status_code == 200

    rows = client.get("/api/bills/search", params={"name": "GREEN TEA"}).json()
    assert len(rows) == 1


def test_ingest_rejects_missing_bill_no(client):
    payload = {"results": [make_result(bill_no=None, source_filename="blurry.jpg")]}
    response = client.post("/api/bills/ingest", json=payload)
    assert response.status_code == 422
    assert "blurry.jpg" in response.json()["detail"]


def test_search_is_case_insensitive_and_filters_on_final_price(client):
    client.post("/api/bills/ingest", json={"results": [make_result()]})

    # SQLite's LIKE ignored case; Postgres's doesn't, hence ILIKE in db.py.
    assert len(client.get("/api/bills/search", params={"name": "green tea"}).json()) == 1
    assert len(client.get("/api/bills/search", params={"name": "GrEeN"}).json()) == 1

    assert client.get("/api/bills/search", params={"min_final_price": 700}).json() == []
    assert len(client.get("/api/bills/search", params={"min_final_price": 600}).json()) == 1
    assert len(client.get("/api/bills/search", params={"max_final_price": 700}).json()) == 1


def test_base_price_is_admin_only(client):
    client.post("/api/bills/ingest", json={"results": [make_result()]})

    public = client.get("/api/bills/search", params={"name": "GREEN"}).json()
    assert "rate" not in public[0], "base price must be absent entirely, not null"
    assert public[0]["final_price"] == 670.0

    wrong = client.get("/api/bills/search", headers={"X-Admin-Password": "nope"})
    assert wrong.status_code == 401

    forbidden = client.get("/api/bills/search", params={"min_base_price": 100})
    assert forbidden.status_code == 403

    allowed = client.get("/api/bills/search", params={"name": "GREEN"}, headers=ADMIN).json()
    assert allowed[0]["rate"] == 567.0

    filtered = client.get("/api/bills/search", params={"min_base_price": 600}, headers=ADMIN)
    assert filtered.json() == []


def test_workbook_groups_by_company_and_404s_for_unknown_supplier(client):
    articles = [
        {
            "company": "JAI MATA DI SRT",
            "product": "GREEN TEA",
            "pcs": 4,
            "rate": 567.0,
            "tax_pct": 10.0,
            "margin_pct": 20.0,
            "final_price": 748.0,
        },
        {
            "company": "ROHAN FAB SRT",
            "product": "MILK CAKE",
            "pcs": 20,
            "rate": 615.0,
            "tax_pct": None,
            "margin_pct": None,
            "final_price": 700.0,
        },
    ]
    client.post("/api/bills/ingest", json={"results": [make_result(articles=articles)]})

    response = client.get("/api/bills/workbook", params={"supplier": "Dindayal Jalan"})
    assert response.status_code == 200

    workbook = load_workbook(io.BytesIO(response.content))
    assert set(workbook.sheetnames) == {"JAI MATA DI SRT", "ROHAN FAB SRT"}

    sheet = workbook["JAI MATA DI SRT"]
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0][0] == "DJ-1"
    assert rows[1] == ("JAI MATA DI SRT", "pc", "price", "tax", "margin", "SP")
    # tax 10% of 567 = 56.7; margin 20% of the tax-inclusive 623.7 = 124.74
    assert rows[2] == ("GREEN TEA", 4, 567.0, 56.7, 124.74, 748.0)

    missing = client.get("/api/bills/workbook", params={"supplier": "Nobody"})
    assert missing.status_code == 404


def test_workbook_keeps_bills_separate_per_company(client):
    """Two bills for one company must stay two dated blocks, not merge."""
    first = make_result(bill_no="DJ-1", bill_date="2025-08-17")
    second = make_result(
        bill_no="DJ-2",
        bill_date="2025-10-02",
        articles=[
            {
                "company": "JAI MATA DI SRT",
                "product": "DIWALI SPL",
                "pcs": 6,
                "rate": 720.0,
                "tax_pct": None,
                "margin_pct": None,
                "final_price": 900.0,
            }
        ],
    )
    client.post("/api/bills/ingest", json={"results": [first, second]})

    response = client.get("/api/bills/workbook", params={"supplier": "Dindayal Jalan"})
    rows = [
        r for r in load_workbook(io.BytesIO(response.content))["JAI MATA DI SRT"].iter_rows(values_only=True)
    ]
    bill_header_rows = [r[0] for r in rows if r[0] in {"DJ-1", "DJ-2"}]
    assert bill_header_rows == ["DJ-1", "DJ-2"]


# --- The app-access gate ---------------------------------------------------

@pytest.mark.parametrize(
    "method,path,kwargs",
    [
        ("get", "/api/bills/providers", {}),
        ("get", "/api/bills/search", {}),
        ("get", "/api/bills/workbook", {"params": {"supplier": "Dindayal Jalan"}}),
        ("post", "/api/bills/ingest", {"json": {"results": []}}),
    ],
)
def test_every_endpoint_refuses_an_unauthenticated_caller(anonymous_client, method, path, kwargs):
    """Without this gate the whole book of record — every supplier, product,
    bill number and selling price — was readable by anyone with the URL, and
    anyone could overwrite it through /ingest."""
    assert getattr(anonymous_client, method)(path, **kwargs).status_code == 401


def test_a_wrong_app_password_is_rejected(anonymous_client):
    response = anonymous_client.get("/api/bills/search", headers={"X-App-Password": "nope"})
    assert response.status_code == 401


def test_health_stays_open_so_the_platform_can_probe_it(anonymous_client):
    assert anonymous_client.get("/health").status_code == 200


def test_search_is_capped_and_pageable(client):
    """A search with no filters used to return every line ever ingested."""
    from app import db

    articles = [
        {"company": "MILL", "product": f"DESIGN {n}", "pcs": 1, "rate": 100.0 + n,
         "tax_pct": None, "margin_pct": None, "final_price": 200.0 + n}
        for n in range(5)
    ]
    client.post("/api/bills/ingest", json={"results": [make_result(articles=articles)]})

    assert len(client.get("/api/bills/search", params={"limit": 2}).json()) == 2
    page_one = client.get("/api/bills/search", params={"limit": 2, "offset": 0}).json()
    page_two = client.get("/api/bills/search", params={"limit": 2, "offset": 2}).json()
    assert page_one != page_two
    assert client.get("/api/bills/search").status_code == 200
    # Over the hard ceiling is rejected by FastAPI's own validation.
    assert client.get(
        "/api/bills/search", params={"limit": db.MAX_SEARCH_LIMIT + 1}
    ).status_code == 422
