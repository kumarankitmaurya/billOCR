"""What staff can and can't see, and what goes to the owner to check.

Staff scan bills and set selling prices; the bill rate is the shop's cost and
must not reach them — not on screen, and not in a response body either. Tax
and margin are withheld too, since applied to a selling price they give the
cost straight back.
"""

import pytest
from fastapi import HTTPException
from openpyxl import Workbook

from app.models import Article, BillExtraction, ExtractionResult
from app.routers.bills import DraftEdit, DraftLine, _admin_view, _apply_edit, _staff_view
from app.services import pricing
from app.services.excel_export import _write_company_sheet
from app.services.ocr_strategy import _apply_pricing_defaults, _staff_flags

COST_KEYS = {"rate", "amount", "tax_pct", "margin_pct"}


def line(product="MILK CAKE", pcs=5, rate=1792.0, amount=8960.0, **kw):
    return Article(company="ROHAN FAB SRT", product=product, pcs=pcs, rate=rate, amount=amount, **kw)


def scan(articles, check_lines=None):
    bill = BillExtraction(supplier="S", bill_no="B1", bill_date="2026-01-01", articles=articles)
    _apply_pricing_defaults(bill)
    return ExtractionResult(
        source_filename="bill.jpg", engine="groq", bill=bill, check_lines=check_lines or []
    )


def keys_in(value) -> set[str]:
    """Every dict key anywhere inside a JSON-shaped value."""
    if isinstance(value, dict):
        return set(value) | {k for v in value.values() for k in keys_in(v)}
    if isinstance(value, list):
        return {k for v in value for k in keys_in(v)}
    return set()


def edit(rows, draft_id="d"):
    return DraftEdit(draft_id=draft_id, bill_no="B1", bill_date="2026-01-01", articles=rows)


def row(ref=0, pcs=5, final_price=2210.0, **kw):
    return DraftLine(ref=ref, company="ROHAN FAB SRT", product="MILK CAKE", pcs=pcs,
                     final_price=final_price, **kw)


# --- Per company / product margins ----------------------------------------

RULES = {("ROHAN FAB SRT", "MILK CAKE"): 25.0, ("ROHAN FAB SRT", ""): 20.0}


def test_a_product_rule_beats_a_company_rule_beats_the_tier():
    assert pricing.margin_for(1000.0, "ROHAN FAB SRT", "MILK CAKE", RULES) == 25.0
    assert pricing.margin_for(1000.0, "ROHAN FAB SRT", "OTHER", RULES) == 20.0
    assert pricing.margin_for(1000.0, "ANOTHER MILL", "MILK CAKE", RULES) == 15.0
    assert pricing.margin_for(2000.0, "ANOTHER MILL", "MILK CAKE", RULES) == 17.0


def test_rules_match_ignoring_case_and_spacing():
    """Names come from OCR; "Rohan Fab  srt" is the same mill."""
    assert pricing.margin_for(1000.0, " Rohan Fab  srt", "milk  cake", RULES) == 25.0


def test_a_scanned_line_is_priced_with_its_rule():
    bill = BillExtraction(supplier="S", bill_no="B", bill_date="2026-01-01", articles=[line(rate=1000.0)])
    _apply_pricing_defaults(bill, RULES)
    article = bill.articles[0]
    assert article.margin_pct == 25.0
    assert article.final_price == float(pricing.sell_price(1000.0, 5.0, 25.0))


# --- The staff view carries no cost ---------------------------------------

def test_staff_view_has_no_cost_fields_anywhere():
    view = _staff_view(scan([line(), line("SAGAR", rate=595.0, amount=2975.0)]), "d1")
    assert not keys_in(view) & COST_KEYS
    assert view["draft_id"] == "d1"
    assert [a["ref"] for a in view["bill"]["articles"]] == [0, 1]
    assert view["bill"]["articles"][0]["final_price"] is not None


def test_admin_view_keeps_everything_and_adds_refs():
    view = _admin_view(scan([line()]), "d1")
    article = view["bill"]["articles"][0]
    assert {"rate", "tax_pct", "margin_pct", "ref"} <= set(article)
    assert "staff_flags" not in view


def test_staff_flags_name_the_line_but_not_the_price():
    """The admin flag reads "5 x 1792 = 8960, bill says 7295" — that is the cost."""
    bill = BillExtraction(supplier="S", bill_no="B", bill_date="2026-01-01",
                          articles=[line(amount=7295.0)])
    flags = _staff_flags(bill)
    assert any("MILK CAKE" in f for f in flags)
    assert not any(n in f for f in flags for n in ("1792", "7295", "8960"))


def test_staff_policy_is_only_the_price_step():
    assert set(pricing.staff_policy()) == {"price_step"}


# --- Applying staff's review to the stored scan ---------------------------

def test_staff_cannot_set_rate_tax_or_margin():
    draft = scan([line()])
    _, _, lines = _apply_edit(draft, edit([row(rate=1.0, tax_pct=0.0, margin_pct=99.0)]), is_admin=False)
    assert (lines[0].rate, lines[0].tax_pct, lines[0].margin_pct) == (1792.0, 5.0, 17.0)


def test_the_owner_can():
    draft = scan([line()])
    _, _, lines = _apply_edit(draft, edit([row(rate=1700.0, margin_pct=20.0)]), is_admin=True)
    assert (lines[0].rate, lines[0].margin_pct) == (1700.0, 20.0)
    assert lines[0].needs_check is False


def test_a_clean_line_priced_above_cost_needs_no_check():
    _, _, lines = _apply_edit(scan([line()]), edit([row()]), is_admin=False)
    assert lines[0].needs_check is False


def test_a_mismatched_line_goes_to_the_owner():
    _, _, lines = _apply_edit(scan([line(amount=7295.0)]), edit([row()]), is_admin=False)
    assert lines[0].needs_check is True


def test_fixing_pcs_clears_the_mismatch():
    """OCR read 4 pcs; the bill says 5. Staff can see and fix pcs, and once
    pcs x rate matches the printed amount there is nothing left to check."""
    draft = scan([line(pcs=4, amount=8960.0)])
    _, _, lines = _apply_edit(draft, edit([row(pcs=5)]), is_admin=False)
    assert lines[0].needs_check is False


def test_no_printed_amount_goes_to_the_owner():
    _, _, lines = _apply_edit(scan([line(amount=None)]), edit([row()]), is_admin=False)
    assert lines[0].needs_check is True


def test_a_whole_bill_misread_goes_to_the_owner():
    _, _, lines = _apply_edit(scan([line()], check_lines=[0]), edit([row()]), is_admin=False)
    assert lines[0].needs_check is True


def test_below_cost_goes_to_the_owner():
    """Staff aren't told why — that would tell them roughly what the cost is."""
    _, _, lines = _apply_edit(scan([line()]), edit([row(final_price=1800.0)]), is_admin=False)
    assert lines[0].needs_check is True


def test_a_hand_added_row_has_no_rate_and_goes_to_the_owner():
    _, _, lines = _apply_edit(scan([line()]), edit([row(), row(ref=None)]), is_admin=False)
    assert lines[1].rate is None
    assert lines[1].needs_check is True


def test_a_ref_outside_the_scan_is_refused():
    with pytest.raises(HTTPException) as exc:
        _apply_edit(scan([line()]), edit([row(ref=3)]), is_admin=False)
    assert exc.value.status_code == 422


def test_staff_edits_to_bill_header_are_kept():
    _, bill, _ = _apply_edit(
        scan([line()]),
        DraftEdit(draft_id="d", bill_no="B9", bill_date="2026-02-02", articles=[row()]),
        is_admin=False,
    )
    assert (bill.bill_no, bill.bill_date) == ("B9", "2026-02-02")


# --- The staff workbook ---------------------------------------------------

def test_staff_workbook_has_no_price_tax_or_margin_columns():
    sheet = Workbook().active
    blocks = [("B1", "2026-01-01", [("MILK CAKE", 5, 1792.0, 2210.0, 17.0, 5.0)])]
    _write_company_sheet(sheet, "ROHAN FAB SRT", blocks, include_cost=False)
    rows = [[c for c in r if c is not None] for r in sheet.iter_rows(values_only=True)]
    assert rows[1] == ["ROHAN FAB SRT", "pc", "SP"]
    assert rows[2] == ["MILK CAKE", 5, 2210.0]


def test_owner_workbook_tolerates_a_line_with_no_rate():
    sheet = Workbook().active
    blocks = [("B1", "2026-01-01", [("NEW ROW", 2, None, 900.0, None, None)])]
    _write_company_sheet(sheet, "ROHAN FAB SRT", blocks, include_cost=True)
    assert list(sheet.iter_rows(values_only=True))[2] == ("NEW ROW", 2, None, None, None, 900.0)
